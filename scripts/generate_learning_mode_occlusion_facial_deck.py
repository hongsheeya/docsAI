#!/usr/bin/env python3
"""Generate a polished FallAI presentation deck from local training artifacts."""

from __future__ import annotations

import datetime as dt
import json
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


PROJECT = Path("/opt/app/project/main")
OUT = PROJECT / "docs" / "presentation"
STORAGE = Path("/opt/app/storage/training/fall-detection")

W, H = 13.333, 7.5

INK = (18, 25, 38)
MUTED = (92, 101, 117)
LINE = (218, 225, 235)
PAPER = (255, 255, 255)
CANVAS = (246, 248, 251)
NAVY = (14, 24, 38)
TEAL = (20, 184, 166)
CYAN = (14, 165, 233)
BLUE = (37, 99, 235)
AMBER = (245, 158, 11)
ROSE = (225, 29, 72)
VIOLET = (124, 58, 237)
GREEN = (34, 197, 94)


def read_json(path: Path) -> dict:
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except Exception:
        return {}


def metric(summary: dict, *keys: str, default=0.0) -> float:
    sources = [summary or {}, (summary or {}).get("best_metrics") or {}]
    for source in sources:
        for key in keys:
            value = source.get(key)
            if value is None:
                continue
            try:
                return float(value)
            except Exception:
                pass
    return float(default)


def pct(value, digits: int = 1) -> str:
    try:
        return f"{float(value) * 100:.{digits}f}%"
    except Exception:
        return "-"


def fmt_int(value, fallback="-") -> str:
    try:
        return f"{int(value):,}"
    except Exception:
        return fallback


def set_bg(slide, color=PAPER):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(*color)


def add_text(
    slide,
    x,
    y,
    w,
    h,
    text,
    *,
    size=18,
    bold=False,
    color=INK,
    align=None,
    valign=MSO_ANCHOR.TOP,
    name="Malgun Gothic",
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.04)
    frame.margin_right = Inches(0.04)
    frame.margin_top = Inches(0.02)
    frame.margin_bottom = Inches(0.02)
    frame.vertical_anchor = valign
    p = frame.paragraphs[0]
    if align is not None:
        p.alignment = align
    run = p.add_run()
    run.text = str(text)
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(*color)
    return box


def add_rect(slide, x, y, w, h, fill, line=None, radius=False):
    shape_type = MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if radius else MSO_AUTO_SHAPE_TYPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(*fill)
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = RGBColor(*line)
    return shape


def add_title(slide, title, kicker="", *, dark=False):
    color = PAPER if dark else INK
    muted = (177, 186, 201) if dark else MUTED
    add_text(slide, 0.65, 0.34, 8.7, 0.3, kicker, size=10, bold=True, color=muted)
    add_text(slide, 0.62, 0.66, 10.9, 0.58, title, size=25, bold=True, color=color)
    add_rect(slide, 0.65, 1.27, 1.1, 0.04, CYAN if dark else TEAL)
    add_rect(slide, 1.82, 1.27, 0.48, 0.04, TEAL if dark else CYAN)


def add_footer(slide, index):
    add_text(slide, 0.65, 7.1, 7.0, 0.22, "FallAI 학습 운영 고도화", size=8, color=(130, 140, 154))
    add_text(slide, 12.1, 7.1, 0.58, 0.22, f"{index:02d}", size=8, color=(130, 140, 154), align=PP_ALIGN.RIGHT)


def new_slide(prs, title, kicker="", idx=0, *, dark=False):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, NAVY if dark else PAPER)
    add_title(slide, title, kicker, dark=dark)
    if idx:
        add_footer(slide, idx)
    return slide


def card(slide, x, y, w, h, title, value="", note="", accent=CYAN, *, dark=False):
    fill = (24, 36, 54) if dark else (248, 250, 252)
    line = (50, 65, 88) if dark else LINE
    title_color = (177, 186, 201) if dark else MUTED
    value_color = PAPER if dark else INK
    note_color = (155, 166, 184) if dark else MUTED
    add_rect(slide, x, y, w, h, fill, line, radius=True)
    add_rect(slide, x, y, 0.08, h, accent)
    add_text(slide, x + 0.2, y + 0.15, w - 0.32, 0.25, title, size=10, bold=True, color=title_color)
    if value:
        add_text(slide, x + 0.2, y + 0.45, w - 0.32, 0.47, value, size=19, bold=True, color=value_color)
    if note:
        add_text(slide, x + 0.2, y + 0.97, w - 0.32, h - 1.05, note, size=9, color=note_color)


def bullet_list(slide, x, y, w, h, items, *, size=15, color=INK, gap=7):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.02)
    frame.margin_right = Inches(0.02)
    for idx, item in enumerate(items):
        p = frame.paragraphs[0] if idx == 0 else frame.add_paragraph()
        p.text = str(item)
        p.font.name = "Malgun Gothic"
        p.font.size = Pt(size)
        p.font.color.rgb = RGBColor(*color)
        p.space_after = Pt(gap)
        p.level = 0
    return box


def status_chip(slide, x, y, label, color, text=""):
    add_rect(slide, x, y, 0.14, 0.14, color, radius=True)
    add_text(slide, x + 0.22, y - 0.04, 1.9, 0.24, label, size=9, bold=True, color=INK)
    if text:
        add_text(slide, x + 0.22, y + 0.17, 2.1, 0.32, text, size=8, color=MUTED)


def section_band(slide, x, y, w, label, body, color):
    add_rect(slide, x, y, w, 0.72, (248, 250, 252), LINE, radius=True)
    add_rect(slide, x, y, 0.09, 0.72, color)
    add_text(slide, x + 0.22, y + 0.12, 2.7, 0.24, label, size=10, bold=True, color=color)
    add_text(slide, x + 2.25, y + 0.1, w - 2.45, 0.42, body, size=13, color=INK)


def range_bar(slide, x, y, w, title, lo, hi, color, note):
    add_text(slide, x, y, 2.6, 0.24, title, size=11, bold=True, color=INK)
    add_rect(slide, x + 2.2, y + 0.08, w, 0.12, (226, 232, 240), radius=True)
    start = max(0, min(1, lo))
    end = max(start, min(1, hi))
    add_rect(slide, x + 2.2 + w * start, y + 0.05, max(0.08, w * (end - start)), 0.18, color, radius=True)
    add_text(slide, x + 2.2 + w + 0.25, y - 0.02, 1.25, 0.22, f"{lo:.2f}-{hi:.2f}", size=9, bold=True, color=color)
    add_text(slide, x, y + 0.32, 10.8, 0.28, note, size=9, color=MUTED)


def interval_bar(slide, x, y, w, label, lo, hi, threshold, color):
    add_text(slide, x, y - 0.02, 1.3, 0.25, label, size=11, bold=True, color=INK)
    add_rect(slide, x + 1.35, y + 0.05, w, 0.15, (226, 232, 240), radius=True)
    start = max(0, min(1, lo))
    end = max(start, min(1, hi))
    add_rect(slide, x + 1.35 + w * start, y + 0.02, max(0.06, w * (end - start)), 0.21, color, radius=True)
    tx = x + 1.35 + w * max(0, min(1, threshold))
    add_rect(slide, tx, y - 0.07, 0.02, 0.39, INK)
    add_text(slide, x + 1.35 + w + 0.2, y - 0.03, 2.0, 0.25, f"{lo:.3f}-{hi:.3f}", size=9, color=MUTED)


def write_notes(path: Path, generated_at: str, pptx_path: Path, slides: list[tuple[str, list[str]]]):
    lines = [
        "# FallAI 발표 원고",
        "",
        f"- 생성: {generated_at}",
        f"- PPTX: `{pptx_path}`",
        "",
        "## 발표 메시지",
        "",
        "- 학습 모드는 추가 확보 데이터를 사이트에서 직접 업로드하고 라벨링한 뒤 성능을 확인해 운영에 반영하는 기능이다.",
        "- 가림 보조는 하체와 좌우 가림을 랜덤 강도로 학습해 자세 불확실성을 줄이는 보조 모델이다.",
        "- 보조 신호는 낙상 주 판단을 대체하지 않고, 조건을 만족할 때만 제한적으로 반영한다.",
        "",
    ]
    for idx, (title, bullets) in enumerate(slides, start=1):
        lines.append(f"## {idx}. {title}")
        lines.append("")
        for bullet in bullets:
            lines.append(f"- {bullet}")
        lines.append("")
    path.write_text("\n".join(lines), encoding="utf-8")


def main() -> int:
    OUT.mkdir(parents=True, exist_ok=True)

    rf = read_json(STORAGE / "rf-fall-v2" / "training_summary.json")
    xg = read_json(STORAGE / "xg-posture" / "training_summary.json")
    occ = read_json(STORAGE / "xg-posture-occlusion-aux" / "training_summary.json")
    face173 = read_json(STORAGE / "facial-state" / "aihub173_driver_state_summary.json")
    status173 = read_json(PROJECT / "outputs" / "continuous_training" / "aihub173_status.json")
    latest_job_id_path = STORAGE / "training-jobs" / "latest_job_id.txt"
    latest_job_id = latest_job_id_path.read_text(encoding="utf-8").strip() if latest_job_id_path.exists() else ""
    latest_job = read_json(STORAGE / "training-jobs" / f"{latest_job_id}.json") if latest_job_id else {}

    def rf_summary_from_job(job: dict) -> dict:
        rf_summary = ((job.get("result") or {}).get("rf_pipeline_training") or {})
        if rf_summary:
            return rf_summary
        for step in job.get("steps") or []:
            if step.get("name") == "rf_pipeline":
                return ((step.get("result") or {}).get("summary") or {})
        return {}

    applied_jobs: list[tuple[float, dict]] = []
    for job_path in (STORAGE / "training-jobs").glob("*.json"):
        job = read_json(job_path)
        if job.get("status") != "applied":
            continue
        if str(job.get("target_model") or "") not in ("rf-fall-v2", "rf-dual"):
            continue
        applied_jobs.append((job_path.stat().st_mtime, job))
    applied_jobs.sort(key=lambda item: item[0], reverse=True)
    operating_job = applied_jobs[0][1] if applied_jobs else latest_job
    operating_rf = rf_summary_from_job(operating_job)
    operating_cv = operating_rf.get("cv") or {}
    operating_dist = operating_rf.get("class_distribution") or operating_job.get("intake_summary") or {}
    operating_upload_y = int(operating_dist.get("Y", 0) or 0)
    operating_upload_n = int(operating_dist.get("N", 0) or 0)
    operating_upload_total = int(operating_rf.get("training_samples", 0) or (operating_job.get("intake_summary") or {}).get("total", 0) or 0)
    operating_job_status = str(operating_job.get("status") or "")
    operating_job_version = str(operating_job.get("version_badge") or "v4")

    latest_rf = ((latest_job.get("result") or {}).get("rf_pipeline_training") or {})
    if not latest_rf:
        for step in latest_job.get("steps") or []:
            if step.get("name") == "rf_pipeline":
                latest_rf = ((step.get("result") or {}).get("summary") or {})
                break
    latest_intake = latest_job.get("intake_summary") or {}
    latest_dist = latest_rf.get("class_distribution") or latest_intake
    upload_y = int(latest_dist.get("Y", 0) or 0)
    upload_n = int(latest_dist.get("N", 0) or 0)
    upload_total = int(latest_rf.get("training_samples", 0) or latest_intake.get("total", 0) or 0)
    upload_need_n = max(0, upload_y - upload_n)
    latest_cv = latest_rf.get("cv") or {}
    latest_rf_f1 = float(latest_cv.get("f1", 0.0) or 0.0)
    latest_rf_acc = float(latest_cv.get("accuracy", 0.0) or 0.0)
    latest_rf_recall = float(latest_cv.get("recall", 0.0) or 0.0)
    latest_job_status = str(latest_job.get("status") or "")

    active173 = float(status173.get("active_macro_f1") or metric(face173, "macro_f1", "f1_macro") or 0.0)
    active173_acc = float(status173.get("active_accuracy") or face173.get("accuracy") or 0.0)

    generated_at = dt.datetime.now(dt.timezone.utc).strftime("%Y-%m-%d %H:%M UTC")
    pptx_path = OUT / "2026-06-10-fallai-learning-mode-occlusion.pptx"
    notes_path = OUT / "2026-06-10-fallai-learning-mode-occlusion.md"

    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)
    notes: list[tuple[str, list[str]]] = []

    # 1. Cover
    s = new_slide(prs, "FallAI 학습 운영 및 보조모델 고도화", "학습 현황 검증 · 운영 모델 · 가림 보조 · 데이터 요구사항", 1, dark=True)
    operating_rf_f1 = float(operating_cv.get("f1", 0.0) or 0.0) if operating_job_status == "applied" else 0.0
    operating_rf_acc = float(operating_cv.get("accuracy", 0.0) or 0.0) if operating_job_status == "applied" else 0.0
    operating_rf_recall = float(operating_cv.get("recall", 0.0) or 0.0) if operating_job_status == "applied" else 0.0
    if operating_rf_f1 <= 0:
        operating_rf_f1 = metric(rf, "f1")
    if operating_rf_acc <= 0:
        operating_rf_acc = metric(rf, "accuracy")
    if operating_rf_recall <= 0:
        operating_rf_recall = metric(rf, "recall")
    operating_rf_name = f"RF-Fall {operating_job_version}" if operating_job_status == "applied" else "RF-Fall v2"
    operating_rf_note = f"Y {fmt_int(operating_upload_y)} / N {fmt_int(operating_upload_n)} 학습 반영" if operating_job_status == "applied" else "낙상 주 판정"

    add_text(s, 0.7, 1.75, 7.2, 1.45, "사이트에서 데이터를 업로드하고,\n학습 현황과 성능을 확인한 뒤,\n운영 모델에 바로 반영하는 구조", size=28, bold=True, color=PAPER)
    add_text(s, 0.73, 3.46, 7.3, 0.56, "발표의 초점은 새 기능 나열이 아니라 운영 신뢰성입니다.", size=15, color=(196, 205, 220))
    card(s, 0.8, 4.7, 2.7, 1.15, operating_rf_name, pct(operating_rf_f1, 1), operating_rf_note, BLUE, dark=True)
    card(s, 3.75, 4.7, 2.7, 1.15, "XG-Posture", pct(metric(xg, "macro_f1", "f1_macro"), 1), "자세 macro F1", TEAL, dark=True)
    card(s, 6.7, 4.7, 2.7, 1.15, "가림 보조", pct(metric(occ, "macro_f1", "f1_macro"), 1), "조건부 자세 보조", VIOLET, dark=True)
    card(s, 9.65, 4.7, 2.7, 1.15, "상태 보조", f"{active173:.1%}", "운영 active", AMBER, dark=True)
    add_text(s, 10.05, 6.7, 2.3, 0.24, generated_at, size=9, color=(177, 186, 201), align=PP_ALIGN.RIGHT)
    notes.append(("표지", ["발표 제목에서 내부 작업 표현을 제거하고 운영 고도화 메시지로 시작한다.", "핵심 지표는 네 개만 보여준다."]))

    # 2. Agenda
    s = new_slide(prs, "목차", "발표 흐름", 2)
    agenda = [
        ("01", "학습 모드", "학습현황 표시와 성능 확인 방식"),
        ("02", "가림 보조모델", "하체·좌우 가림 학습 설계와 낙상 판단 영향"),
        ("03", "운영 모델 검증", "보조 신호를 제한적으로 반영하는 방식"),
        ("04", "데이터 요구사항", "필요 데이터, 레이블, 수량, 우선순위"),
    ]
    for i, (num, title, body) in enumerate(agenda):
        y = 1.65 + i * 1.05
        add_text(s, 0.95, y, 0.75, 0.42, num, size=20, bold=True, color=CYAN)
        add_text(s, 1.85, y, 2.1, 0.32, title, size=17, bold=True, color=INK)
        add_text(s, 4.2, y + 0.03, 7.1, 0.32, body, size=14, color=MUTED)
        add_rect(s, 0.95, y + 0.63, 10.8, 0.02, LINE)
    notes.append(("목차", ["표지 다음에 네 개 섹션으로 발표 구조를 명확히 한다."]))

    # 3. Learning mode visibility
    s = new_slide(prs, "학습 모드: 현황 표시 개선", "진행 중만 보던 화면에서 검증 가능한 상태판으로", 3)
    section_band(s, 0.8, 1.55, 11.8, "만든 이유", "추가로 확보한 낙상/비낙상 영상을 사이트에서 직접 업로드하고, Y/N·자세 라벨과 함께 intake에 저장하기 위해 만들었다.", CYAN)
    section_band(s, 0.8, 2.65, 11.8, "운영 흐름", "업로드 등록 → 라벨 검증 → 백그라운드 학습 job → 성능 확인 → 운영 모델 적용 순서로 진행한다.", TEAL)
    section_band(s, 0.8, 3.75, 11.8, "표시 개선", "진행 중뿐 아니라 완료, 차단, 실패 상태를 함께 보여서 학습 가능 여부와 적용 여부를 분리한다.", BLUE)
    status_chip(s, 1.0, 5.35, "학습 중", GREEN, "프로세스·로그 표시")
    status_chip(s, 3.35, 5.35, "완료", BLUE, "성능 지표 표시")
    status_chip(s, 5.7, 5.35, "데이터 대기", AMBER, "원천/라벨 부족")
    status_chip(s, 8.05, 5.35, "중단", ROSE, "PID 종료·실패 로그")
    notes.append(("학습 모드", ["학습 모드의 목적은 추가 확보 데이터를 사이트에서 직접 등록하고 학습 결과를 운영에 반영하는 것이다.", "학습 가능 조건과 성능 지표를 화면에서 확인한다."]))

    # 4. Active state auxiliary model
    s = new_slide(prs, "상태 보조 운영 모델", "현재 서비스에 사용 중인 active 기준", 4)
    card(s, 0.85, 1.6, 3.45, 1.35, "운영 모델", "Active", "상태 보조 모델 사용 중", BLUE)
    card(s, 4.9, 1.6, 3.45, 1.35, "Macro F1", f"{active173:.4f}", "운영 기준 성능", TEAL)
    card(s, 8.95, 1.6, 3.45, 1.35, "Accuracy", f"{active173_acc:.4f}", "운영 기준 성능", CYAN)
    section_band(s, 0.95, 3.75, 11.15, "역할", "상태 보조 모델은 졸림, 하품, 주의저하처럼 낙상 판단에 참고할 수 있는 상태 신호를 제공한다.", TEAL)
    section_band(s, 0.95, 4.75, 11.15, "운영 방식", "낙상 여부를 단독으로 결정하지 않고 RF-Fall과 자세 판단 결과를 보조하는 제한 신호로 사용한다.", BLUE)
    notes.append(("상태 보조 운영 모델", ["현재 운영 중인 active 상태 보조 모델은 macro F1 0.9054, accuracy 0.9020이다.", "발표자료에는 현재 서비스에 쓰는 운영 정보만 넣는다."]))

    # 5. Current model state
    s = new_slide(prs, "현재 모델 운영 상태", "서비스에 적용된 모델 기준", 5)
    card(s, 0.75, 1.55, 2.75, 1.35, operating_rf_name, "적용 완료", f"CV F1 {pct(operating_rf_f1, 1)}", BLUE)
    card(s, 3.75, 1.55, 2.75, 1.35, "XG-Posture", "목표 충족", f"Macro F1 {pct(metric(xg, 'macro_f1', 'f1_macro'), 1)}", TEAL)
    card(s, 6.75, 1.55, 2.75, 1.35, "가림 보조", "학습 완료", f"Macro F1 {pct(metric(occ, 'macro_f1', 'f1_macro'), 1)}", VIOLET)
    card(s, 9.75, 1.55, 2.75, 1.35, "상태 보조", "운영 active", f"F1 {active173:.4f}", CYAN)
    section_band(s, 0.95, 3.75, 11.15, "운영 모델 4개", "RF-Fall, XG-Posture, 가림 보조, 상태 보조를 각각 목적별 active 모델로 운영한다.", BLUE)
    section_band(s, 0.95, 4.75, 11.15, "RF-Fall 적용 결과", f"사이트 업로드 데이터 Y {fmt_int(operating_upload_y)} / N {fmt_int(operating_upload_n)} / 총 {fmt_int(operating_upload_total)}개로 학습했고, CV F1 {pct(operating_rf_f1, 1)} · Acc {pct(operating_rf_acc, 1)} · Recall {pct(operating_rf_recall, 1)}를 확인한 뒤 적용했다.", TEAL)
    section_band(s, 0.95, 5.75, 11.15, "이후 운영", "추가 데이터가 들어오면 모델별 라벨 조건을 맞춘 뒤 학습하고, 현재 운영 모델과 비교해 더 좋은 모델만 반영한다.", VIOLET)
    notes.append(("현재 모델 운영 상태", [f"{operating_rf_name}은 적용 완료 상태이며 CV F1 {operating_rf_f1:.4f}이다.", "오늘 발표에서는 운영 핵심 모델 4개만 표시한다."]))

    # 6. Occlusion design
    s = new_slide(prs, "가림 보조모델 학습 설계", "하체 가림과 좌우 가림을 따로 만들고 자세 모델의 약점을 보완", 6)
    card(s, 0.8, 1.55, 3.55, 1.45, "하체 가림", "y축 기준", "무릎·발·하체가 사라지는 상황", VIOLET)
    card(s, 4.85, 1.55, 3.55, 1.45, "좌우 가림", "x축 기준", "문틀·가구·프레임 밖으로 한쪽 몸이 가려짐", CYAN)
    card(s, 8.9, 1.55, 3.55, 1.45, "랜덤 강도", "mild / moderate / severe", "고정 가림이 아니라 범위 내 샘플링", AMBER)
    section_band(s, 0.95, 3.95, 11.15, "학습 목적", "가림 자체를 낙상으로 판단하는 것이 아니라, 가림 때문에 posture가 sit/lie/stand 사이에서 흔들리는 문제를 줄인다.", TEAL)
    section_band(s, 0.95, 4.95, 11.15, "운영 방식", "평소에는 기존 자세 모델을 쓰고, keypoint confidence나 visibility가 낮은 경우에만 보조 확률을 blend한다.", BLUE)
    notes.append(("가림 학습 설계", ["하체와 좌우 가림을 분리해서 hard-case로 학습한다.", "가림 보조는 조건부로만 동작한다."]))

    # 7. Occlusion ranges
    s = new_slide(prs, "가림 랜덤 수치", "좌우 가림도 고정값이 아니라 범위 샘플링으로 구성", 7)
    range_bar(s, 0.85, 1.55, 5.65, "하체 mild", 0.68, 0.78, VIOLET, "하체 일부만 가려지는 약한 케이스")
    range_bar(s, 0.85, 2.35, 5.65, "하체 moderate", 0.55, 0.68, VIOLET, "허벅지 아래가 상당 부분 가려지는 케이스")
    range_bar(s, 0.85, 3.15, 5.65, "하체 severe", 0.42, 0.55, VIOLET, "하체 대부분이 사라지는 강한 케이스")
    range_bar(s, 0.85, 4.25, 5.65, "좌/우 mild", 0.18, 0.30, CYAN, "화면 한쪽 일부가 가려지는 케이스")
    range_bar(s, 0.85, 5.05, 5.65, "좌/우 moderate", 0.30, 0.45, CYAN, "옆면 가림이 자세 판정에 영향을 주는 케이스")
    range_bar(s, 0.85, 5.85, 5.65, "좌/우 severe", 0.45, 0.58, CYAN, "몸의 절반 가까이가 가려지는 강한 케이스")
    notes.append(("가림 랜덤 수치", ["하체는 0.42에서 0.78 사이, 좌우는 0.18에서 0.58 사이를 강도별로 샘플링한다.", "좌우 가림도 랜덤화했다."]))

    # 8. Occlusion impact
    s = new_slide(prs, "가림이 낙상 판단에 미치는 영향", "최종 판단을 대체하지 않고 불확실한 posture를 안정화", 8)
    add_text(s, 0.9, 1.65, 2.0, 0.5, "입력 pose", size=18, bold=True, color=INK, align=PP_ALIGN.CENTER)
    add_rect(s, 3.0, 1.82, 1.1, 0.05, LINE)
    add_text(s, 4.3, 1.65, 2.5, 0.5, "가림 위험 계산", size=18, bold=True, color=INK, align=PP_ALIGN.CENTER)
    add_rect(s, 7.0, 1.82, 1.1, 0.05, LINE)
    add_text(s, 8.25, 1.65, 2.7, 0.5, "조건부 blend", size=18, bold=True, color=INK, align=PP_ALIGN.CENTER)
    add_rect(s, 0.95, 2.3, 10.8, 0.02, LINE)
    card(s, 0.8, 3.05, 3.55, 1.4, "트리거", "confidence < 0.38", "또는 lower visibility < 0.42", AMBER)
    card(s, 4.85, 3.05, 3.55, 1.4, "자세 보정", "margin < 0.07", "주 모델이 애매할 때만 blend", CYAN)
    card(s, 8.9, 3.05, 3.55, 1.4, "낙상 RF", "occlusion_fall_risk", "65개 feature 중 하나로 반영", VIOLET)
    section_band(s, 0.95, 5.25, 11.15, "검증 방향", "가림이 많은 구간에서 sit/lie/stand confusion이 줄어드는지와 RF-Fall false positive가 늘지 않는지를 함께 본다.", TEAL)
    notes.append(("가림 영향", ["가림 보조는 직접 낙상 판정기가 아니다.", "자세 안정화와 occlusion_fall_risk feature를 통해 간접 반영된다."]))

    # 9. Conservative decision policy
    s = new_slide(prs, "운영 판단 보수화", "보조 신호는 주 판단을 대체하지 않도록 제한", 9)
    section_band(s, 0.85, 1.55, 11.65, "주 판단", "낙상 여부는 RF-Fall과 자세 흐름을 기준으로 판단하고, 보조 모델은 특정 조건에서만 참고한다.", BLUE)
    section_band(s, 0.85, 2.55, 11.65, "가림 조건", "keypoint confidence, visibility, posture margin이 낮은 경우에만 가림 보조를 blend한다.", CYAN)
    section_band(s, 0.85, 3.55, 11.65, "상태 조건", "상태 보조는 졸림·주의저하처럼 위험 해석에 필요한 맥락만 제공하고 단독 낙상 판정에는 쓰지 않는다.", TEAL)
    section_band(s, 0.85, 4.55, 11.65, "적용 원칙", "새 후보 모델은 현재 운영 모델보다 지표가 낮으면 자동 적용하지 않고, 후보 로그로만 남긴다.", AMBER)
    notes.append(("운영 판단 보수화", ["보조 모델은 주 판단을 대체하지 않는다.", "현재 운영 모델보다 낮은 후보는 적용하지 않는다."]))

    # 10. Data and labels
    s = new_slide(prs, "교수님께 요청할 데이터", "사용자가 직접 확보하기 어려운 통제 촬영 셋", 10)
    section_band(s, 0.75, 1.55, 11.85, "Hard negative N", f"현재 추가 업로드는 Y {fmt_int(upload_y)} / N {fmt_int(upload_n)}이다. 1:1 균형을 위해 우선 N {fmt_int(upload_need_n)}개 안팎이 더 필요하다.", BLUE)
    section_band(s, 0.75, 2.45, 11.85, "비낙상 장면", "앉기, 눕기, 바닥 작업, 침대/소파 이동, 물건 줍기, 운동·스트레칭처럼 낙상과 비슷하지만 실제 N인 장면.", TEAL)
    section_band(s, 0.75, 3.35, 11.85, "경계 낙상 Y", "천천히 주저앉음, 미끄러짐, 의자·침대 낙상, 바로 회복하는 짧은 낙상처럼 recall이 흔들리는 장면.", ROSE)
    section_band(s, 0.75, 4.25, 11.85, "가림/상태 통제", "하체·좌우 가림 비율, 문틀·가구 가림, 졸림·주의저하·정상 상태를 함께 기록한다.", VIOLET)
    section_band(s, 0.75, 5.35, 11.85, "필수 레이블", "Y/N, 자세, 낙상 시작/종료 시간, 가림 방향·비율, 상태 라벨, 카메라 각도, 조명, 장소를 한 클립 단위로 저장한다.", AMBER)
    notes.append(("교수님께 요청할 데이터", ["사용자가 직접 만들기 어려운 통제 촬영 hard-case가 필요하다.", f"현재 추가 업로드는 Y {upload_y}, N {upload_n}이므로 N 균형 데이터가 최우선이다."]))

    # 11. Data count targets
    s = new_slide(prs, "데이터 수와 수집 우선순위", "현재 업로드와 교수님 요청 목표", 11)
    card(s, 0.75, 1.55, 2.85, 1.25, "현재 업로드", f"Y {fmt_int(upload_y)} / N {fmt_int(upload_n)}", f"총 {fmt_int(upload_total)}개", BLUE)
    card(s, 3.9, 1.55, 2.85, 1.25, "RF 운영 모델", pct(operating_rf_f1, 1), f"Acc {pct(operating_rf_acc, 1)} · 적용 완료", TEAL)
    card(s, 7.05, 1.55, 2.85, 1.25, "N 추가 목표", f"+{fmt_int(upload_need_n)}", "1:1 균형 우선", ROSE)
    card(s, 10.2, 1.55, 2.85, 1.25, "검증셋", "Y/N 200+200", "학습 미사용 holdout", AMBER)
    section_band(s, 0.95, 3.7, 11.15, "1순위", f"비낙상 hard negative N {fmt_int(upload_need_n)}개 내외: 바닥에 앉기·눕기, 기대기, 물건 줍기, 침대/소파 이동.", ROSE)
    section_band(s, 0.95, 4.55, 11.15, "2순위", "가림 paired set 약 360개: 하체/좌/우 × mild/moderate/severe × Y/N을 같은 동작 조건으로 촬영.", CYAN)
    section_band(s, 0.95, 5.4, 11.15, "3순위", "경계 낙상 Y 300~400개와 상태 보조 300~500개: 졸림·주의저하·정상 상태를 분리해 라벨링.", TEAL)
    section_band(s, 0.95, 6.25, 11.15, "검증 원칙", "교수님 제작 데이터 일부는 학습에 넣지 않고 holdout으로 남겨 실제 성능을 다시 표시한다.", BLUE)
    notes.append(("데이터 수와 우선순위", ["현재 RF 운영 모델은 적용 완료됐지만 Y가 더 많아 N hard negative가 계속 필요하다.", "교수님께는 통제 촬영, 가림 paired set, 독립 holdout을 요청한다."]))

    # 12. Conclusion
    s = new_slide(prs, "결론", "보이는 학습, 적용된 운영 모델, 제한된 보조 신호", 12)
    card(s, 0.85, 1.75, 3.6, 1.6, "학습 모드", "상태 표시 강화", "완료·중단·데이터 대기까지 보이게 한다.", CYAN)
    card(s, 4.85, 1.75, 3.6, 1.6, "가림 보조", "하체+좌우 랜덤", "가림 때문에 흔들리는 자세 판단을 조건부로 보정한다.", VIOLET)
    card(s, 8.85, 1.75, 3.6, 1.6, "운영 비교", "낮은 후보 차단", "현재 운영 모델보다 낮은 결과는 자동 적용하지 않는다.", AMBER)
    add_text(s, 1.1, 4.55, 11.1, 0.7, "발표 메시지", size=17, bold=True, color=INK, align=PP_ALIGN.CENTER)
    add_text(s, 1.2, 5.1, 10.9, 0.72, "FallAI는 사이트 업로드 데이터로 모델을 학습하고,\n성능이 확인된 모델을 운영에 반영하며,\n가림·상태 보조 신호를 보수적으로 통제한다.", size=20, bold=True, color=INK, align=PP_ALIGN.CENTER)
    notes.append(("결론", ["학습 운영의 투명성과 운영 반영 결과를 최종 메시지로 닫는다.", "가림과 상태 보조 신호는 보수적으로 운용한다."]))

    prs.save(pptx_path)
    write_notes(notes_path, generated_at, pptx_path, notes)
    print(pptx_path)
    print(notes_path)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
