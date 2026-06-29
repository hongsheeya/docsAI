#!/usr/bin/env python3
from __future__ import annotations

import datetime as dt
import json
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


PROJECT = Path("/mnt/data/wiz/project/main")
STORAGE = Path("/mnt/data/wiz/storage/training/fall-detection")
OUT = PROJECT / "docs" / "presentation"
PPTX = OUT / "2026-06-15-fallai-system-update.pptx"
NOTES = OUT / "2026-06-15-fallai-system-update.md"

BG = RGBColor(248, 250, 252)
INK = RGBColor(15, 23, 42)
MUTED = RGBColor(100, 116, 139)
LINE = RGBColor(226, 232, 240)
GREEN = RGBColor(16, 185, 129)
BLUE = RGBColor(37, 99, 235)
INDIGO = RGBColor(79, 70, 229)
AMBER = RGBColor(245, 158, 11)
ROSE = RGBColor(225, 29, 72)
WHITE = RGBColor(255, 255, 255)


def read_json(path: Path) -> dict:
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except Exception:
        return {}


def metric(summary: dict, *keys: str) -> float:
    sources = [
        summary,
        summary.get("best_metrics") or {},
        summary.get("group_cv") or {},
        summary.get("sequence_group_cv") or {},
        summary.get("validation") or {},
        summary.get("operational_validation") or {},
    ]
    for source in sources:
        for key in keys:
            try:
                value = source.get(key)
                if value is not None:
                    return float(value)
            except Exception:
                pass
    return 0.0


def pct(value: float, digits: int = 1) -> str:
    return f"{float(value) * 100:.{digits}f}%" if value else "-"


def tx(shape, text: str, size: int = 18, bold: bool = False, color=INK, align=None):
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    if align is not None:
        p.alignment = align
    r = p.runs[0] if p.runs else p.add_run()
    r.font.name = "Malgun Gothic"
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    return shape


def add_text(slide, x, y, w, h, text, size=18, bold=False, color=INK, align=None):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    return tx(shape, text, size=size, bold=bold, color=color, align=align)


def fill(shape, color, line=LINE):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.color.rgb = line
    shape.line.width = Pt(1)


def add_card(slide, x, y, w, h, title, body="", accent=BLUE, value=None):
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    fill(shape, WHITE, LINE)
    try:
        shape.adjustments[0] = 0.08
    except Exception:
        pass
    bar = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(x), Inches(y), Inches(0.07), Inches(h))
    fill(bar, accent, accent)
    add_text(slide, x + 0.22, y + 0.16, w - 0.35, 0.28, title, size=11, bold=True, color=MUTED)
    if value is not None:
        add_text(slide, x + 0.22, y + 0.48, w - 0.35, 0.45, value, size=23, bold=True, color=INK)
        if body:
            add_text(slide, x + 0.22, y + 0.98, w - 0.35, h - 1.08, body, size=10, color=MUTED)
    else:
        add_text(slide, x + 0.22, y + 0.50, w - 0.35, h - 0.60, body, size=12, color=INK)
    return shape


def add_bullets(slide, x, y, w, h, items, size=16, color=INK):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    for idx, item in enumerate(items):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = str(item)
        p.level = 0
        p.font.name = "Malgun Gothic"
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.space_after = Pt(8)
    return shape


def add_header(slide, title, subtitle=""):
    add_text(slide, 0.55, 0.34, 8.7, 0.42, title, size=22, bold=True, color=INK)
    if subtitle:
        add_text(slide, 0.58, 0.78, 8.7, 0.25, subtitle, size=10, color=MUTED)
    line = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0.55), Inches(1.10), Inches(12.25), Inches(0.01))
    fill(line, LINE, LINE)


def add_footer(slide, page):
    add_text(slide, 11.7, 7.05, 1.0, 0.2, str(page), size=8, color=MUTED, align=PP_ALIGN.RIGHT)


def new_slide(prs, title, subtitle="", page=0):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = BG
    add_header(slide, title, subtitle)
    add_footer(slide, page)
    return slide


def add_bar(slide, x, y, w, label, value, color=BLUE):
    add_text(slide, x, y, 2.1, 0.25, label, size=11, bold=True, color=INK)
    bg = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(x + 2.25), Inches(y + 0.04), Inches(w), Inches(0.16))
    fill(bg, RGBColor(226, 232, 240), RGBColor(226, 232, 240))
    fg = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(x + 2.25), Inches(y + 0.04), Inches(max(0.05, w * min(1.0, value))), Inches(0.16))
    fill(fg, color, color)
    add_text(slide, x + 2.25 + w + 0.18, y - 0.01, 0.85, 0.24, pct(value), size=10, bold=True, color=INK)


def main() -> int:
    OUT.mkdir(parents=True, exist_ok=True)
    generated = dt.datetime.now(dt.timezone.utc).strftime("%Y-%m-%d %H:%M UTC")

    rf = read_json(STORAGE / "rf-fall-v2" / "training_summary.json")
    rf_pipe = read_json(STORAGE / "rf-pipeline" / "training_summary.json")
    xg = read_json(STORAGE / "xg-posture" / "training_summary.json")
    occ = read_json(STORAGE / "xg-posture-occlusion-aux" / "training_summary.json")
    face82 = read_json(STORAGE / "facial-state" / "aihub82_facial_emotion_summary.json")
    face173 = read_json(STORAGE / "facial-state" / "aihub173_driver_state_summary.json")
    dl = read_json(PROJECT / "outputs" / "continuous_training" / "xg-posture_status.json")

    rf_f1 = metric(rf, "f1", "macro_f1", "f1_macro")
    rf_pipe_f1 = metric(rf_pipe, "f1", "macro_f1", "f1_macro")
    xg_f1 = metric(xg, "f1", "macro_f1", "f1_macro")
    occ_f1 = metric(occ, "f1", "macro_f1", "f1_macro")
    face82_f1 = metric(face82, "macro_f1", "f1_macro", "f1")
    face173_f1 = metric(face173, "macro_f1", "f1_macro", "f1")

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    notes: list[tuple[str, list[str]]] = []

    # 1 Cover
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(12, 18, 32)
    add_text(slide, 0.75, 0.85, 8.7, 0.55, "FallAI 시스템 개선 및 운영 현황", size=28, bold=True, color=WHITE)
    add_text(slide, 0.78, 1.55, 8.5, 0.34, "온디바이스 방향 · 모델 업로드/관리 · 성능 비교 · 커스텀 최종 모델 · 가림보조 개선", size=14, color=RGBColor(203, 213, 225))
    add_card(slide, 0.82, 5.35, 3.1, 0.95, "운영 낙상 모델", "RF-Dual/RF-Pipeline 기준", GREEN, pct(max(rf_f1, rf_pipe_f1)),)
    add_card(slide, 4.15, 5.35, 3.1, 0.95, "자세 모델", "XG-Posture active", BLUE, pct(xg_f1))
    add_card(slide, 7.48, 5.35, 3.1, 0.95, "가림보조", "90% 목표 개선 진행", AMBER, pct(occ_f1))
    add_text(slide, 0.78, 6.85, 4.0, 0.24, f"생성: {generated}", size=9, color=RGBColor(148, 163, 184))
    notes.append(("표지", ["발표 범위: 온디바이스 방향, 모델 관리 기능, 비교 기능, 커스텀 최종 모델, 가림보조 개선."]))

    # 2 Agenda
    slide = new_slide(prs, "발표 흐름", "저번 발표 이후 실제로 바뀐 기능과 모델 운영 구조 중심", 2)
    agenda = [
        "현재 운영 모델과 목표",
        "학습 모드와 모델 업로드 기능",
        "모델 관리/비교/삭제 유예",
        "커스텀 최종 모델 구성",
        "가림보조 모델 개선 방향",
        "온디바이스 적용 방향",
        "필요 데이터와 향후 일정",
    ]
    for i, item in enumerate(agenda):
        x = 0.8 + (i % 2) * 5.9
        y = 1.45 + (i // 2) * 1.05
        add_card(slide, x, y, 5.35, 0.72, f"{i + 1}. {item}", "", BLUE if i < 4 else GREEN)
    notes.append(("발표 흐름", agenda))

    # 3 Current objective
    slide = new_slide(prs, "현재 목표", "최종 목표는 운영 전체 모델 97%, 단기 우선순위는 가림보조 90% 이상", 3)
    add_card(slide, 0.75, 1.45, 3.65, 1.35, "현재 운영 기준", "운영 active 모델은 유지하면서 후보가 더 좋을 때만 적용합니다.", GREEN, "안정성 우선")
    add_card(slide, 4.85, 1.45, 3.65, 1.35, "단기 병목", "가림보조는 pseudo-label만으로는 90% 확신이 어렵습니다.", AMBER, "실제 라벨 필요")
    add_card(slide, 8.95, 1.45, 3.65, 1.35, "운영 방식", "세부 모델을 선택해 커스텀 최종 모델로 묶어 실시간/업로드에 적용합니다.", INDIGO, "조합 운영")
    add_bullets(slide, 0.9, 3.25, 11.6, 2.3, [
        "무작정 데이터 대기만 하지 않고, 사용 가능한 데이터로 먼저 학습/검증합니다.",
        "성능이 낮은 후보는 active를 덮어쓰지 않고 후보/로그로 남깁니다.",
        "후보 모델은 삭제 예약으로 정리하되, 운영 모델과 현재 참조 파일은 보호합니다.",
    ], size=16)
    notes.append(("현재 목표", ["전체 장기 목표는 97%, 가림보조 단기 목표는 90%."]))

    # 4 Learning mode
    slide = new_slide(prs, "학습 모드를 만든 이유", "추가 확보 데이터를 사이트에서 직접 업로드하고 학습까지 연결하기 위한 구조", 4)
    steps = [
        ("1. 데이터 등록", "영상/압축 파일 업로드, fall/non-fall 및 자세 라벨 저장"),
        ("2. 학습 실행", "등록 데이터와 AI-Hub 데이터를 합쳐 백그라운드 학습"),
        ("3. 검증 표시", "F1/Acc/Recall, ETA, 로그, 실패 원인을 학습 모드에 표시"),
        ("4. 적용 제어", "active보다 낮으면 자동 적용하지 않고 후보로 보존"),
    ]
    for i, (title, body) in enumerate(steps):
        add_card(slide, 0.8 + (i % 2) * 6.0, 1.45 + (i // 2) * 1.75, 5.35, 1.18, title, body, [BLUE, GREEN, INDIGO, AMBER][i])
    notes.append(("학습 모드", [body for _, body in steps]))

    # 5 Model upload
    slide = new_slide(prs, "모델 업로드와 관리", "파일/폴더 업로드로 외부 학습 모델도 registry에 등록", 5)
    add_card(slide, 0.85, 1.45, 3.55, 1.35, "업로드", "모델 파일과 summary JSON을 함께 올리면 버전/성능을 자동 표시", BLUE, "파일/폴더")
    add_card(slide, 4.85, 1.45, 3.55, 1.35, "관리", "active, 후보, 업로드, 커스텀 최종 모델을 한 목록에서 관리", INDIGO, "Registry")
    add_card(slide, 8.85, 1.45, 3.55, 1.35, "삭제", "삭제 예약 후 24시간 유예, 유예 중 취소 가능", ROSE, "1일 유예")
    add_bullets(slide, 0.95, 3.45, 11.4, 1.75, [
        "후보 보호는 해제했습니다. 운영 active와 현재 참조 중인 파일만 삭제를 차단합니다.",
        "불필요 후보는 자동 정리로 family별 최고 후보만 남기고 유예 삭제할 수 있습니다.",
        "삭제 유예는 실수로 지웠을 때 복구할 시간을 주기 위한 안전장치입니다.",
    ], size=15)
    notes.append(("모델 업로드/관리", ["파일/폴더 업로드, registry, 삭제 유예, 삭제 취소."]))

    # 6 Comparison
    slide = new_slide(prs, "모델 성능 비교", "선택한 버전만 비교해서 운영 후보를 빠르게 확인", 6)
    add_bar(slide, 0.9, 1.65, 6.8, "RF-Fall", rf_f1, GREEN)
    add_bar(slide, 0.9, 2.25, 6.8, "RF-Pipeline", rf_pipe_f1, GREEN)
    add_bar(slide, 0.9, 2.85, 6.8, "XG-Posture", xg_f1, BLUE)
    add_bar(slide, 0.9, 3.45, 6.8, "가림 보조", occ_f1, AMBER)
    add_bar(slide, 0.9, 4.05, 6.8, "AI-Hub 173 상태", face173_f1, INDIGO)
    add_bar(slide, 0.9, 4.65, 6.8, "AI-Hub 82 표정", face82_f1, ROSE)
    add_card(slide, 9.15, 1.65, 3.1, 2.1, "비교 화면", "선택한 모델만 그래프에 표시하고, 버전/성능/샘플 수를 함께 보여줍니다.", BLUE)
    add_card(slide, 9.15, 4.05, 3.1, 1.25, "판단 기준", "F1, Accuracy, Recall, Precision 기준 전환", GREEN)
    notes.append(("모델 비교", ["선택한 버전만 비교, 대표 선택, 지표 전환."]))

    # 7 Custom final model
    slide = new_slide(prs, "커스텀 최종 모델", "세부 모델을 하나씩 골라 하나의 운영 조합으로 저장", 7)
    families = [
        ("RF-Fall", "낙상 이진 판단"),
        ("RF-Pipeline", "fallback/보조 RF"),
        ("XG-Posture", "자세·행동 분류"),
        ("가림 보조", "하체/좌우 가림 조건부 보정"),
        ("AI-Hub 173 상태", "졸림·하품·주의저하"),
        ("AI-Hub 82 표정", "얼굴 표정 보조"),
    ]
    for i, (name, desc) in enumerate(families):
        add_card(slide, 0.75 + (i % 3) * 4.15, 1.45 + (i // 3) * 1.38, 3.55, 0.94, name, desc, [GREEN, GREEN, BLUE, AMBER, INDIGO, ROSE][i])
    add_bullets(slide, 0.95, 4.45, 11.4, 1.15, [
        "선택한 세부 모델 조합은 저장된 최종 모델로 남고, 필요할 때 다시 적용할 수 있습니다.",
        "최종 모델을 적용하면 실시간 분석과 업로드 분석 모두 같은 세부 모델 조합을 사용합니다.",
    ], size=15)
    notes.append(("커스텀 최종 모델", ["세부 모델을 선택해 하나의 최종 운영 조합으로 저장/적용."]))

    # 8 Occlusion
    slide = new_slide(prs, "가림보조 모델", "0.38은 몸의 38% 가림이 아니라 keypoint 평균 confidence 기준", 8)
    add_card(slide, 0.8, 1.4, 3.25, 1.2, "avg_conf < 0.38", "keypoint 평균 confidence가 낮을 때 보조 후보", BLUE, "0.38")
    add_card(slide, 4.25, 1.4, 3.25, 1.2, "lower visibility < 0.42", "하체 keypoint가 충분히 보이지 않는 상황", AMBER, "0.42")
    add_card(slide, 7.7, 1.4, 3.25, 1.2, "margin < 0.07", "주 자세 모델이 sit/lie/stand에서 애매한 상황", INDIGO, "0.07")
    add_bullets(slide, 0.9, 3.25, 11.6, 1.9, [
        "학습 증강은 하체 가림과 좌/우 가림을 mild, moderate, severe 범위로 랜덤 생성합니다.",
        "현재 active 가림보조는 macro F1 86%대라 운영 보조로는 가능하지만 90% 확신은 부족합니다.",
        "90% 이상을 확실히 노리려면 pseudo-label이 아니라 실제 자세 라벨이 포함된 가림 데이터가 필요합니다.",
    ], size=15)
    notes.append(("가림보조", ["0.38/0.42/0.07 의미, 랜덤 가림, 실제 라벨 필요."]))

    # 9 Data acquisition
    slide = new_slide(prs, "71461 우선 재수신과 학습 순서", "가림보조 90% 달성 전까지 다른 대형 학습은 중단하고 71461을 우선 처리", 9)
    eta = str(dl.get("eta_text") or "다운로드 상태 확인 중")
    latest = str(dl.get("latest_log") or "")
    add_card(slide, 0.8, 1.45, 5.55, 1.35, "현재 다운로드", latest[:130], BLUE, eta[:34])
    add_card(slide, 6.85, 1.45, 5.55, 1.35, "학습 순서", "71461 라벨 감지 → 가림보조 우선 학습 → 90% 달성 후 다른 모델 순차 재개", GREEN, "가림 우선")
    add_bullets(slide, 0.95, 3.45, 11.3, 1.8, [
        "persistent 영역 용량 부족 위험 때문에 71461은 임시 대용량 루트에서 먼저 받습니다.",
        "라벨 JSON이 풀리면 watcher가 자동으로 가림보조 학습을 시작합니다.",
        "0.90 미만 후보는 active를 덮어쓰지 않고 결과와 로그만 남깁니다.",
    ], size=15)
    notes.append(("71461 우선 처리", [eta, latest]))

    # 10 Data request
    slide = new_slide(prs, "교수님께 요청할 데이터", "우리가 직접 확보하기 어려운 실제 라벨 데이터", 10)
    add_card(slide, 0.8, 1.35, 3.65, 1.55, "실제 가림 자세 라벨", "하체/좌/우 가림, mild/moderate/severe, stand/sit/lie/walk/run/fall", AMBER, "최우선")
    add_card(slide, 4.85, 1.35, 3.65, 1.55, "near-miss/비낙상", "앉기, 눕기, 숙이기, 침대/의자 전이, 카메라 사각", BLUE, "오탐 감소")
    add_card(slide, 8.9, 1.35, 3.65, 1.55, "정상 표정/평온", "무표정, 평온, 차분, 얼굴 미검출/부분 검출", GREEN, "표정 안정화")
    add_bullets(slide, 0.95, 3.45, 11.3, 1.85, [
        "권장 단위: 4~6초 clip, class별 최소 200~300개, 사람/장소가 겹치지 않는 split 정보 포함.",
        "필수 라벨: fall 여부, 자세 class, 가림 방향/정도, 카메라 각도, 얼굴 검출 가능 여부.",
        "이 데이터가 있어야 pseudo-label 편향 없이 가림보조 90% 이상을 검증할 수 있습니다.",
    ], size=15)
    notes.append(("필요 데이터", ["실제 가림 자세 라벨, near-miss, 정상 표정 데이터."]))

    # 11 On-device
    slide = new_slide(prs, "온디바이스 적용 방향", "현재 프로그램을 그대로 넣기보다 추론 패키지와 관리 서버를 분리하는 방식이 현실적", 11)
    add_card(slide, 0.8, 1.4, 3.55, 1.5, "그대로 내장", "브라우저 UI, Python 학습, AI-Hub 관리, 대용량 모델/데이터까지 포함하면 장치 부담이 큽니다.", ROSE, "비권장")
    add_card(slide, 4.85, 1.4, 3.55, 1.5, "추천 구조", "장치에는 추론 런타임만, 학습/모델 관리는 서버에서 수행합니다.", GREEN, "Hybrid")
    add_card(slide, 8.9, 1.4, 3.55, 1.5, "배포 단위", "선택한 커스텀 최종 모델 조합을 export해서 장치에 탑재합니다.", BLUE, "Bundle")
    add_bullets(slide, 0.95, 3.65, 11.25, 1.5, [
        "장치: 카메라 입력, pose 추론, RF/XGBoost/보조 모델 추론, 결과 표시.",
        "서버: 데이터 업로드, 학습, 모델 비교, 커스텀 최종 모델 생성, 원격 업데이트.",
    ], size=15)
    notes.append(("온디바이스", ["그대로 내장은 어렵고, 추론 패키지+관리 서버 분리가 현실적."]))

    # 12 Reliability
    slide = new_slide(prs, "운영 안정성 장치", "성능 낮은 후보와 실수 삭제가 운영 모델을 망가뜨리지 않게 설계", 12)
    add_card(slide, 0.8, 1.45, 3.6, 1.25, "자동 적용 제한", "active보다 낮은 후보는 운영 모델을 덮어쓰지 않음", GREEN, "보호")
    add_card(slide, 4.85, 1.45, 3.6, 1.25, "삭제 유예", "삭제 예약 후 24시간 동안 취소 가능", ROSE, "복구")
    add_card(slide, 8.9, 1.45, 3.6, 1.25, "버전 기록", "모델 버전, feature 수치, threshold, summary를 registry에 보존", BLUE, "추적")
    add_bullets(slide, 0.95, 3.35, 11.25, 1.75, [
        "모델 목록의 성능 수치는 summary JSON 기준으로 표시해 비교 화면과 맞춥니다.",
        "현재 사용 중인 커스텀 최종 모델은 삭제 차단, 미사용 후보는 삭제 예약 가능.",
        "백그라운드 학습 완료 시 성능/적용 여부가 학습 상태와 모델 목록에 남습니다.",
    ], size=15)
    notes.append(("운영 안정성", ["자동 적용 제한, 삭제 유예, 버전 기록."]))

    # 13 Timeline
    slide = new_slide(prs, "남은 작업", "발표 이후 구현/검증 우선순위", 13)
    add_bullets(slide, 0.9, 1.45, 11.4, 4.7, [
        "71461 라벨 확보 후 가림보조 90% 후보 학습 및 holdout 검증",
        "커스텀 최종 모델 export 기능 추가: 장치 배포용 bundle 생성",
        "모델 비교 화면에 feature/threshold diff 추가",
        "학습 모드에서 pause/resume/retry를 버튼으로 제어",
        "삭제 유예 모델 복구, 정리 이력, 모델 적용 이력 표시 강화",
    ], size=18)
    notes.append(("남은 작업", ["가림보조 90%, export, diff, pause/resume/retry."]))

    # 14 Conclusion
    slide = new_slide(prs, "결론", "현재 운영 가능한 모델을 유지하면서, 가림보조와 배포 구조를 개선 중", 14)
    add_card(slide, 0.8, 1.45, 3.55, 1.45, "운영", "현재 active 모델은 유지 가능하고, 더 좋은 후보만 검증 후 적용합니다.", GREEN, "사용 가능")
    add_card(slide, 4.85, 1.45, 3.55, 1.45, "기능", "업로드, 비교, 모델 관리, 커스텀 최종 모델 생성 흐름을 추가했습니다.", BLUE, "관리 가능")
    add_card(slide, 8.9, 1.45, 3.55, 1.45, "개선", "가림보조 90%는 실제 자세 라벨 가림 데이터가 핵심입니다.", AMBER, "데이터 필요")
    add_bullets(slide, 0.95, 3.75, 11.25, 1.3, [
        "다음 단계는 71461/실제 가림 라벨을 이용해 가림보조 모델을 90% 이상으로 올리는 것입니다.",
        "온디바이스는 전체 시스템 내장이 아니라 추론 bundle 중심으로 진행하는 것이 현실적입니다.",
    ], size=16)
    notes.append(("결론", ["운영 가능, 관리 기능 추가, 가림보조 데이터 필요, 온디바이스는 bundle 중심."]))

    prs.save(PPTX)
    lines = [f"# FallAI 시스템 개선 및 운영 현황", "", f"- 생성: {generated}", f"- PPTX: `{PPTX}`", ""]
    for idx, (title, bullets) in enumerate(notes, start=1):
        lines.append(f"## {idx}. {title}")
        for bullet in bullets:
            lines.append(f"- {bullet}")
        lines.append("")
    NOTES.write_text("\n".join(lines), encoding="utf-8")
    print(PPTX)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
